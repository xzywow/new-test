import json
import tempfile
import unittest
from pathlib import Path
from decimal import Decimal
from openpyxl import Workbook, load_workbook
from import_rules import COLUMNS, EXPECTED_HEADERS, SkuRegistry, selling_price_usd, validate_record, export_import

DEMO = Path(__file__).parents[1] / 'examples/import-records.demo.json'


class ImportTests(unittest.TestCase):
    def test_required_and_positive(self):
        record = json.loads(DEMO.read_text(encoding='utf-8'))[0]
        self.assertEqual(validate_record(record, allow_missing_sku=True), [])
        for key in ['申报金额', '商品净重', '英文报关名']:
            changed = dict(record)
            changed.pop(key)
            self.assertTrue(validate_record(changed, allow_missing_sku=True))
        for value in [0, -1, 'NaN', 'Infinity', True]:
            self.assertTrue(validate_record({**record, '长': value}, allow_missing_sku=True))

    def test_sku_persists_and_resets_each_day(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'sku.sqlite3'
            registry = SkuRegistry(path)
            self.assertEqual(registry.assign('a', '20261015'), '20261015-0001')
            self.assertEqual(registry.assign('b', '20261015'), '20261015-0002')
            registry.close()
            registry = SkuRegistry(path)
            self.assertEqual(registry.assign('a', '20261016'), '20261015-0001')
            self.assertEqual(registry.assign('c', '20261016'), '20261016-0001')
            registry.close()

    def test_fixed_markup_and_explicit_fx(self):
        self.assertEqual(selling_price_usd('10', '5'), Decimal('15.00'))
        self.assertEqual(selling_price_usd('70', '5', 'CNY', '7'), Decimal('15.00'))
        with self.assertRaises(ValueError):
            selling_price_usd('70', '5', 'CNY')

    def test_template_export_text_and_units(self):
        with tempfile.TemporaryDirectory() as tmp:
            template, output = Path(tmp) / 'template.xlsx', Path(tmp) / 'output.xlsx'
            workbook = Workbook()
            for (_, column), heading in zip(COLUMNS.items(), EXPECTED_HEADERS):
                workbook.active.cell(1, column, heading)
            workbook.active.cell(1, 27, '销售员')
            workbook.save(template)
            workbook.close()
            record = json.loads(DEMO.read_text(encoding='utf-8'))[0]
            record.update(SKU='20261015-0001', 英文名称='=NOT_A_FORMULA()')
            original = template.read_bytes()
            export_import(template, output, [record])
            self.assertEqual(template.read_bytes(), original)
            result = load_workbook(output)
            self.assertEqual(result.active['A2'].value, record['SKU'])
            self.assertEqual(result.active['E2'].data_type, 's')
            self.assertEqual(result.active['H2'].value, 100)
            self.assertEqual(result.active['S2'].value, 5)
            self.assertEqual(result.active['AA1'].value, '销售员')
            result.close()
            with self.assertRaises(ValueError):
                export_import(template, output, [record])
            with self.assertRaises(ValueError):
                export_import(template, Path(tmp) / 'bad.xlsx', [{**record, '中文名称': ''}])
