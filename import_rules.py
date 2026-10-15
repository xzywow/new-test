"""Local product import preparation; no platform writes or inferred physical data."""
import re
import sqlite3
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from urllib.parse import urlsplit
from openpyxl import load_workbook

REQUIRED = ['SKU', '中文名称', '英文名称', '商品净重', '来源URL', '长', '宽', '高',
            '英文报关名', '中文报关名', '申报重量', '申报金额']
NUMERIC = ['商品净重', '长', '宽', '高', '申报重量', '申报金额']
COLUMNS = {'SKU': 1, '中文名称': 4, '英文名称': 5, '商品净重': 8, '长': 11, '宽': 12,
           '高': 13, '来源URL': 14, '英文报关名': 16, '中文报关名': 17, '申报重量': 18, '申报金额': 19}
EXPECTED_HEADERS = ['*SKU', '中文名称', '英文名称', '商品净重', '长', '宽', '高',
                    '来源URL', '英文报关名', '中文报关名', '申报重量', '申报金额']


def positive(value):
    if isinstance(value, bool):
        raise ValueError('必须为正数')
    try:
        number = Decimal(str(value))
    except InvalidOperation as error:
        raise ValueError('必须为正数') from error
    if not number.is_finite() or number <= 0 or number > Decimal('1000000000'):
        raise ValueError('必须为有效正数且不超过十亿')
    return number


def selling_price_usd(cost, markup_usd, currency='USD', cny_per_usd=None):
    cost = positive(cost)
    markup = Decimal(str(markup_usd))
    if not markup.is_finite() or markup < 0:
        raise ValueError('固定加价必须为非负美元金额')
    if currency == 'CNY':
        cost /= positive(cny_per_usd)
    elif currency != 'USD':
        raise ValueError('仅支持 USD/CNY 成本，不能隐式猜测汇率')
    return (cost + markup).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def validate_record(record, allow_missing_sku=False):
    errors = []
    for field in REQUIRED:
        if field == 'SKU' and allow_missing_sku and not record.get(field):
            continue
        value = record.get(field)
        if value is None or not str(value).strip():
            errors.append(f'{field}：必填')
        elif field in NUMERIC:
            try:
                positive(value)
            except ValueError:
                errors.append(f'{field}：必须为有效正数')
        elif not isinstance(value, str):
            errors.append(f'{field}：必须为文本')
    url = record.get('来源URL', '')
    try:
        parsed = urlsplit(url if isinstance(url, str) else '')
        if parsed.scheme not in {'http', 'https'} or not parsed.hostname or parsed.username or parsed.password:
            errors.append('来源URL：必须为不含账号密码的 HTTP(S) 地址')
    except ValueError:
        errors.append('来源URL：地址无效')
    return errors


class SkuRegistry:
    """One key per source variant. Persist this DB when moving/restarting the tool."""
    def __init__(self, path):
        self.connection = sqlite3.connect(path)
        self.connection.execute('CREATE TABLE IF NOT EXISTS sku (record_key TEXT PRIMARY KEY, code TEXT UNIQUE NOT NULL)')
        self.connection.execute('CREATE TABLE IF NOT EXISTS sequence (day TEXT PRIMARY KEY, value INTEGER NOT NULL)')

    def assign(self, record_key, day=None):
        if not isinstance(record_key, str) or not record_key.strip():
            raise ValueError('必须提供稳定的 record_key（来源商品与规格的唯一组合）')
        day = day or date.today().strftime('%Y%m%d')
        if not re.fullmatch(r'\d{8}', day):
            raise ValueError('日期必须为 YYYYMMDD')
        date(int(day[:4]), int(day[4:6]), int(day[6:]))
        try:
            self.connection.execute('BEGIN IMMEDIATE')
            existing = self.connection.execute('SELECT code FROM sku WHERE record_key=?', (record_key,)).fetchone()
            if existing:
                self.connection.commit()
                return existing[0]
            previous = self.connection.execute('SELECT value FROM sequence WHERE day=?', (day,)).fetchone()
            value = previous[0] + 1 if previous else 1
            self.connection.execute('INSERT OR REPLACE INTO sequence VALUES (?, ?)', (day, value))
            code = f'{day}-{value:04d}'
            self.connection.execute('INSERT INTO sku VALUES (?, ?)', (record_key, code))
            self.connection.commit()
            return code
        except BaseException:
            self.connection.rollback()
            raise

    def close(self):
        self.connection.close()


def export_import(template, output, records):
    template, output = Path(template), Path(output)
    if template.resolve() == output.resolve() or output.exists():
        raise ValueError('必须输出到新的文件，不能覆盖模板或已有文件')
    if not records:
        raise ValueError('没有商品，不导出空结果')
    failures = []
    seen = set()
    for index, record in enumerate(records, 1):
        failures.extend(f'第{index}条 {error}' for error in validate_record(record))
        sku = record.get('SKU')
        if isinstance(sku, str):
            if sku in seen:
                failures.append(f'第{index}条 SKU 重复')
            seen.add(sku)
    if failures:
        raise ValueError('\n'.join(failures))
    workbook = load_workbook(template)
    try:
        sheet = workbook.worksheets[0]
        for (field, column), heading in zip(COLUMNS.items(), EXPECTED_HEADERS):
            actual = str(sheet.cell(1, column).value or '').split('\n')[0]
            actual = actual.split('（')[0].strip()
            if actual != heading:
                raise ValueError(f'模板表头不匹配：{field}')
        if any(cell.value is not None for row in sheet.iter_rows(min_row=2) for cell in row):
            raise ValueError('模板含现有商品数据；请使用空白模板，避免覆盖')
        for row_number, record in enumerate(records, 2):
            for field, column in COLUMNS.items():
                cell = sheet.cell(row_number, column)
                if field in NUMERIC:
                    cell.value = float(positive(record[field]))
                else:
                    cell.value = str(record[field]).strip()
                    cell.data_type = 's'
        # Exclusive file creation protects existing outputs, including concurrent runs.
        with output.open('xb') as handle:
            workbook.save(handle)
    finally:
        workbook.close()
